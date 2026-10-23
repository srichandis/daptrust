#!/usr/bin/env python3
"""
Import DAP Educational Trust Master Account Excel data into Laravel SQLite database.
Reads each FY sheet, parses transactions, and inserts into SQLite.
"""

import msoffcrypto
import io
import openpyxl
import sqlite3
import os
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP

# Find the database file
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'database', 'database.sqlite')

# Find Excel file
EXCEL_FILE = None
for f in os.listdir(BASE_DIR):
    if f.endswith('.xlsx') and 'MASTER' in f.upper():
        EXCEL_FILE = os.path.join(BASE_DIR, f)
        break

if not EXCEL_FILE:
    print("ERROR: Could not find Excel file with 'MASTER' in name")
    exit(1)

PASSWORD = 'dap6977'

# Account keys mapping (matches Laravel account IDs)
ACCT_MAP = {'cash': 1, 'ab': 2, 'sbi': 3, 'hdfc': 4}


def load_workbook():
    """Load the encrypted Excel file."""
    print(f"Loading Excel: {EXCEL_FILE}")
    with open(EXCEL_FILE, 'rb') as f:
        ms = msoffcrypto.OfficeFile(f)
        ms.load_key(password=PASSWORD)
        decrypted = io.BytesIO()
        ms.decrypt(decrypted)
        decrypted.seek(0)
        return openpyxl.load_workbook(decrypted, data_only=True)


def parse_date(s):
    """Parse date strings in various formats."""
    if not s or not isinstance(s, str):
        return None
    s = s.strip()
    for fmt in ('%d.%m.%Y', '%d.%m.%y', '%d/%m/%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def safe_float(v):
    """Convert cell value to float."""
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0


def find_header_row(ws):
    """Find the header row by looking for 'DATE' in column B."""
    for r in range(1, min(6, ws.max_row + 1)):
        for cell in ws[r]:
            if cell.column_letter == 'B' and cell.value and str(cell.value).strip().upper() == 'DATE':
                return r
    return 3


def build_col_map(ws):
    """Map column letter -> (account_key, entry_type, is_balance)."""
    HEADER = {
        'CASH DONOR': ('cash', 'donation_income', False),
        'BANK TRANSFER': ('cash', 'bank_transfer', False),
        'STATIONARY CHARGES': ('cash', 'stationary_charges', False),
        'CASH DONATION': ('cash', 'donation_expense', False),
        'CASH BALANCE': ('cash', None, True),
        'AB DONOR': ('ab', 'donation_income', False),
        'AB DONATION': ('ab', 'donation_expense', False),
        'AB BANK CHARGES': ('ab', 'bank_charges', False),
        'AB BANK INTEREST': ('ab', 'bank_interest', False),
        'AB FD INTEREST': ('ab', 'fd_interest', False),
        'AB PROFESSIONAL CHARGES': ('ab', 'professional_charges', False),
        'AB STATIONARY CHARGES': ('ab', 'stationary_charges', False),
        'AB BANK TRANSFER': ('ab', 'bank_transfer', False),
        'AB IT RETURN': ('ab', 'it_return', False),
        'FD - AB': ('ab', 'fd_investment', False),
        'ANDHRA BANK BALANCE': ('ab', None, True),
        'DONOR - SBI': ('sbi', 'donation_income', False),
        'DONATION - SBI': ('sbi', 'donation_expense', False),
        'SBI BANK CHARGES': ('sbi', 'bank_charges', False),
        'SBI BANK INTEREST': ('sbi', 'bank_interest', False),
        'SBI FD INTEREST': ('sbi', 'fd_interest', False),
        'SBI PROFESSIONAL CHARGES': ('sbi', 'professional_charges', False),
        'SBI STATIONARY CHARGES': ('sbi', 'stationary_charges', False),
        'SBI BANK TRANSFER': ('sbi', 'bank_transfer', False),
        'SBI IT RETURN': ('sbi', 'it_return', False),
        'FD - SBI': ('sbi', 'fd_investment', False),
        'SBI BANK BALANCE': ('sbi', None, True),
        'DONOR - HDFC': ('hdfc', 'donation_income', False),
        'DONATION - HDFC': ('hdfc', 'donation_expense', False),
        'HDFC BANK CHARGES': ('hdfc', 'bank_charges', False),
        'HDFC BANK INTEREST': ('hdfc', 'bank_interest', False),
        'HDFC FD INTEREST': ('hdfc', 'fd_interest', False),
        'HDFC PROFESSIONAL CHARGES': ('hdfc', 'professional_charges', False),
        'HDFC STATIONARY CHARGES': ('hdfc', 'stationary_charges', False),
        'HDFC BANK TRANSFER': ('hdfc', 'bank_transfer', False),
        'HDFC IT RETURN': ('hdfc', 'it_return', False),
        'FD - HDFC': ('hdfc', 'fd_investment', False),
        'HDFC BANK BALANCE': ('hdfc', None, True),
        'DONATION TO OPS SAIVA PERAVAI': ('sbi', 'donation_expense', False),
    }

    fd_maturity_count = 0
    rd_maturity_count = 0
    mf_maturity_count = 0

    cols = []
    header_row = find_header_row(ws)

    for cell in ws[header_row]:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        letter = cell.column_letter
        idx = cell.column - 1

        # Handle FD MATURITY AMOUNT (appears multiple times)
        if name == 'FD MATURITY AMOUNT':
            acct = ['ab', 'sbi', 'hdfc'][fd_maturity_count] if fd_maturity_count < 3 else 'ab'
            cols.append((letter, acct, 'fd_maturity', False))
            fd_maturity_count += 1
            continue

        # Handle RD MATURITY AMOUNT (appears multiple times)
        if name == 'RD MATURITY AMOUNT':
            acct = ['ab', 'sbi', 'hdfc'][rd_maturity_count] if rd_maturity_count < 3 else 'ab'
            cols.append((letter, acct, 'rd_maturity', False))
            rd_maturity_count += 1
            continue

        # Handle MF MATURITY AMOUNT (appears multiple times)
        if name == 'MF MATURITY AMOUNT':
            acct = ['ab', 'sbi', 'hdfc'][mf_maturity_count] if mf_maturity_count < 3 else 'ab'
            cols.append((letter, acct, 'mf_maturity', False))
            mf_maturity_count += 1
            continue

        # Handle CASH DONOR variants
        if 'CASH DONOR' in name and 'CASH DONATION' not in name:
            cols.append((letter, 'cash', 'donation_income', False))
            continue

        # Handle RD/MF interest columns
        if name.startswith('AB RD INTEREST'):
            cols.append((letter, 'ab', 'rd_interest', False))
        elif name.startswith('AB MF INTEREST'):
            cols.append((letter, 'ab', 'mf_interest', False))
        elif name.startswith('RD - AB'):
            cols.append((letter, 'ab', 'rd_investment', False))
        elif name.startswith('MF - AB'):
            cols.append((letter, 'ab', 'mf_investment', False))
        elif name.startswith('SBI RD INTEREST'):
            cols.append((letter, 'sbi', 'rd_interest', False))
        elif name.startswith('SBI MF INTEREST'):
            cols.append((letter, 'sbi', 'mf_interest', False))
        elif name.startswith('RD - SBI'):
            cols.append((letter, 'sbi', 'rd_investment', False))
        elif name.startswith('MF - SBI'):
            cols.append((letter, 'sbi', 'mf_investment', False))
        elif name.startswith('HDFC RD INTEREST'):
            cols.append((letter, 'hdfc', 'rd_interest', False))
        elif name.startswith('HDFC MF INTEREST'):
            cols.append((letter, 'hdfc', 'mf_interest', False))
        elif name.startswith('RD - HDFC'):
            cols.append((letter, 'hdfc', 'rd_investment', False))
        elif name.startswith('MF - HDFC'):
            cols.append((letter, 'hdfc', 'mf_investment', False))
        elif name in HEADER:
            acct, etype, bal = HEADER[name]
            cols.append((letter, acct, etype, bal))
        # Skip unknown columns

    return cols


def parse_sheet(ws, fy_name):
    """Parse all data rows from an FY sheet."""
    col_map = build_col_map(ws)
    header_row = find_header_row(ws)
    data_start = header_row + 2

    entries = []
    transfers = []

    for row_idx in range(data_start, ws.max_row + 1):
        row = ws[row_idx]
        date_val = row[1].value if len(row) > 1 else None  # Column B
        desc_val = row[2].value if len(row) > 2 else None  # Column C

        if not date_val:
            continue
        if isinstance(desc_val, str) and 'C/F' in desc_val.upper():
            continue

        entry_date = parse_date(str(date_val))
        if entry_date is None:
            continue

        description = str(desc_val).strip() if desc_val else ''
        row_transfers = []

        for letter, acct, etype, is_bal in col_map:
            if is_bal or etype is None:
                continue
            col_idx = openpyxl.utils.column_index_from_string(letter) - 1
            if col_idx >= len(row):
                continue
            val = safe_float(row[col_idx].value)
            if val == 0:
                continue

            if etype == 'bank_transfer':
                row_transfers.append((entry_date, description, acct, val))
            else:
                donor = description if etype == 'donation_income' else None
                entries.append({
                    'entry_type': etype,
                    'entry_date': entry_date,
                    'financial_year': fy_name,
                    'account': acct,
                    'amount': abs(val),
                    'particulars': description,
                    'donor_name': donor,
                })

        # Handle transfers: pair them up
        used = set()
        for i, (dt, desc, src, amt) in enumerate(row_transfers):
            if i in used:
                continue
            for j, (dt2, desc2, src2, amt2) in enumerate(row_transfers):
                if j in used or j == i:
                    continue
                if src != src2 and amt == amt2:
                    transfers.append({
                        'entry_type': 'bank_transfer',
                        'entry_date': dt,
                        'financial_year': fy_name,
                        'account': src,
                        'to_account': src2,
                        'amount': amt,
                        'particulars': desc,
                    })
                    used.add(i)
                    used.add(j)
                    break
            else:
                transfers.append({
                    'entry_type': 'bank_transfer',
                    'entry_date': dt,
                    'financial_year': fy_name,
                    'account': src,
                    'to_account': None,
                    'amount': amt,
                    'particulars': desc,
                })
                used.add(i)

    return entries, transfers


def main():
    # Check if database exists
    if not os.path.exists(DB_PATH):
        print(f"ERROR: Database not found at {DB_PATH}")
        print("Please run: php artisan migrate")
        exit(1)

    print(f"Database: {DB_PATH}")

    # Load Excel
    wb = load_workbook()

    # Connect to SQLite
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    # Clear existing entries
    print("\nClearing existing entries...")
    cur.execute("DELETE FROM entries")
    cur.execute("DELETE FROM sqlite_sequence WHERE name='entries'")
    conn.commit()
    print("Entries cleared.")

    # Check current accounts
    cur.execute("SELECT id, name FROM accounts ORDER BY id")
    accounts = cur.fetchall()
    print(f"\nExisting accounts ({len(accounts)}):")
    for aid, name in accounts:
        print(f"  ID {aid}: {name}")

    if len(accounts) == 0:
        print("\nERROR: No accounts found. Please run the account seeder first.")
        conn.close()
        exit(1)

    # Parse and import each FY sheet
    total_entries = 0
    total_transfers = 0
    fy_sheets = [s for s in wb.sheetnames if s.startswith('FY')]

    print(f"\nFound {len(fy_sheets)} FY sheets")

    for fy in fy_sheets:
        ws = wb[fy]
        entries, transfers = parse_sheet(ws, fy)

        # Insert entries
        for e in entries:
            cur.execute(
                """INSERT INTO entries 
                   (entry_type, entry_date, financial_year, account_id, amount, 
                    particulars, donor_name, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now'))""",
                (e['entry_type'], str(e['entry_date']), e['financial_year'],
                 ACCT_MAP.get(e['account']), e['amount'],
                 e['particulars'], e.get('donor_name'))
            )

        # Insert transfers
        for t in transfers:
            cur.execute(
                """INSERT INTO entries 
                   (entry_type, entry_date, financial_year, account_id, to_account_id,
                    amount, particulars, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now'))""",
                (t['entry_type'], str(t['entry_date']), t['financial_year'],
                 ACCT_MAP.get(t['account']),
                 ACCT_MAP.get(t['to_account']) if t['to_account'] else None,
                 t['amount'], t['particulars'])
            )

        conn.commit()
        total_entries += len(entries)
        total_transfers += len(transfers)
        print(f"  {fy}: {len(entries)} entries + {len(transfers)} transfers")

    # Verify
    print("\n" + "="*50)
    print("IMPORT COMPLETE")
    print("="*50)

    cur.execute("SELECT COUNT(*) FROM entries")
    count = cur.fetchone()[0]
    print(f"Total entries in database: {count}")

    cur.execute("SELECT entry_type, COUNT(*) as cnt FROM entries GROUP BY entry_type ORDER BY cnt DESC")
    print("\nEntries by type:")
    for etype, cnt in cur.fetchall():
        print(f"  {etype}: {cnt}")

    # Per-account summary
    cur.execute("""
        SELECT a.name, COUNT(e.id) as entry_count
        FROM accounts a
        LEFT JOIN entries e ON e.account_id = a.id
        GROUP BY a.id, a.name
    """)
    print("\nAccount summary:")
    for name, cnt in cur.fetchall():
        print(f"  {name}: {cnt} entries")

    cur.close()
    conn.close()
    print("\nDone!")


if __name__ == '__main__':
    main()
