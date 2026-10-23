#!/usr/bin/env python3
"""
Verify imported data against Excel.
Reads each FY sheet and compares totals per entry type and per account.
"""

import msoffcrypto
import io
import openpyxl
import pymysql
from datetime import datetime
import os

# Find Excel file
FILE_PATH = None
for f in os.listdir(os.path.dirname(os.path.abspath(__file__))):
    if f.endswith('.xlsx') and 'MASTER' in f.upper():
        FILE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), f)
        break

PASSWORD = 'dap6977'

DB_CONFIG = {
    'host': '127.0.0.1',
    'port': 3306,
    'user': 'root',
    'password': 'root',
    'database': 'dap',
    'charset': 'utf8mb4',
}

ACCT_MAP = {'cash': 1, 'ab': 2, 'sbi': 3, 'hdfc': 4}


def load_workbook():
    with open(FILE_PATH, 'rb') as f:
        ms = msoffcrypto.OfficeFile(f)
        ms.load_key(password=PASSWORD)
        d = io.BytesIO()
        ms.decrypt(d)
        d.seek(0)
        return openpyxl.load_workbook(d, data_only=True)


def parse_date(s):
    if not s or not isinstance(s, str):
        return None
    s = s.strip()
    for fmt in ('%d.%m.%Y', '%d.%m.%y', '%d/%m/%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def safe_float(v):
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0


def find_header_row(ws):
    for r in range(1, min(6, ws.max_row + 1)):
        for cell in ws[r]:
            if cell.column_letter == 'B' and cell.value and str(cell.value).strip().upper() == 'DATE':
                return r
    return 3


def build_col_map(ws):
    HEADER = {
        'CASH DONOR': ('cash', 'donation_income'),
        'BANK TRANSFER': ('cash', 'bank_transfer'),
        'STATIONARY CHARGES': ('cash', 'stationary_charges'),
        'CASH DONATION': ('cash', 'donation_expense'),
        'CASH BALANCE': ('cash', None),
        'AB DONOR': ('ab', 'donation_income'),
        'AB DONATION': ('ab', 'donation_expense'),
        'AB BANK CHARGES': ('ab', 'bank_charges'),
        'AB BANK INTEREST': ('ab', 'bank_interest'),
        'AB FD INTEREST': ('ab', 'fd_interest'),
        'AB PROFESSIONAL CHARGES': ('ab', 'professional_charges'),
        'AB STATIONARY CHARGES': ('ab', 'stationary_charges'),
        'AB BANK TRANSFER': ('ab', 'bank_transfer'),
        'AB IT RETURN': ('ab', 'it_return'),
        'FD - AB': ('ab', 'fd_investment'),
        'ANDHRA BANK BALANCE': ('ab', None),
        'DONOR - SBI': ('sbi', 'donation_income'),
        'DONATION - SBI': ('sbi', 'donation_expense'),
        'SBI BANK CHARGES': ('sbi', 'bank_charges'),
        'SBI BANK INTEREST': ('sbi', 'bank_interest'),
        'SBI FD INTEREST': ('sbi', 'fd_interest'),
        'SBI PROFESSIONAL CHARGES': ('sbi', 'professional_charges'),
        'SBI STATIONARY CHARGES': ('sbi', 'stationary_charges'),
        'SBI BANK TRANSFER': ('sbi', 'bank_transfer'),
        'SBI IT RETURN': ('sbi', 'it_return'),
        'FD - SBI': ('sbi', 'fd_investment'),
        'SBI BANK BALANCE': ('sbi', None),
        'DONOR - HDFC': ('hdfc', 'donation_income'),
        'DONATION - HDFC': ('hdfc', 'donation_expense'),
        'HDFC BANK CHARGES': ('hdfc', 'bank_charges'),
        'HDFC BANK INTEREST': ('hdfc', 'bank_interest'),
        'HDFC FD INTEREST': ('hdfc', 'fd_interest'),
        'HDFC PROFESSIONAL CHARGES': ('hdfc', 'professional_charges'),
        'HDFC STATIONARY CHARGES': ('hdfc', 'stationary_charges'),
        'HDFC BANK TRANSFER': ('hdfc', 'bank_transfer'),
        'HDFC IT RETURN': ('hdfc', 'it_return'),
        'FD - HDFC': ('hdfc', 'fd_investment'),
        'HDFC BANK BALANCE': ('hdfc', None),
        'DONATION TO OPS SAIVA PERAVAI': ('sbi', 'donation_expense'),
    }

    fd_maturity_count = 0
    rd_maturity_count = 0
    mf_maturity_count = 0

    cols = []
    header_row = find_header_row(ws)

    for cell in ws[header_row]:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        letter = cell.column_letter

        if name == 'FD MATURITY AMOUNT':
            acct = ['ab', 'sbi', 'hdfc'][fd_maturity_count] if fd_maturity_count < 3 else 'ab'
            cols.append((letter, acct, 'fd_maturity'))
            fd_maturity_count += 1
            continue

        if name == 'RD MATURITY AMOUNT':
            acct = ['ab', 'sbi', 'hdfc'][rd_maturity_count] if rd_maturity_count < 3 else 'ab'
            cols.append((letter, acct, 'rd_maturity'))
            rd_maturity_count += 1
            continue

        if name == 'MF MATURITY AMOUNT':
            acct = ['ab', 'sbi', 'hdfc'][mf_maturity_count] if mf_maturity_count < 3 else 'ab'
            cols.append((letter, acct, 'mf_maturity'))
            mf_maturity_count += 1
            continue

        if 'CASH DONOR' in name and 'CASH DONATION' not in name:
            cols.append((letter, 'cash', 'donation_income'))
            continue

        if name.startswith('AB RD INTEREST'):
            cols.append((letter, 'ab', 'rd_interest'))
        elif name.startswith('AB MF INTEREST'):
            cols.append((letter, 'ab', 'mf_interest'))
        elif name.startswith('RD - AB'):
            cols.append((letter, 'ab', 'rd_investment'))
        elif name.startswith('MF - AB'):
            cols.append((letter, 'ab', 'mf_investment'))
        elif name.startswith('SBI RD INTEREST'):
            cols.append((letter, 'sbi', 'rd_interest'))
        elif name.startswith('SBI MF INTEREST'):
            cols.append((letter, 'sbi', 'mf_interest'))
        elif name.startswith('RD - SBI'):
            cols.append((letter, 'sbi', 'rd_investment'))
        elif name.startswith('MF - SBI'):
            cols.append((letter, 'sbi', 'mf_investment'))
        elif name.startswith('HDFC RD INTEREST'):
            cols.append((letter, 'hdfc', 'rd_interest'))
        elif name.startswith('HDFC MF INTEREST'):
            cols.append((letter, 'hdfc', 'mf_interest'))
        elif name.startswith('RD - HDFC'):
            cols.append((letter, 'hdfc', 'rd_investment'))
        elif name.startswith('MF - HDFC'):
            cols.append((letter, 'hdfc', 'mf_investment'))
        elif name in HEADER:
            acct, etype = HEADER[name]
            cols.append((letter, acct, etype))

    return cols


def parse_sheet_totals(ws, fy_name):
    """Parse totals per entry type and per account from an FY sheet."""
    col_map = build_col_map(ws)
    header_row = find_header_row(ws)
    data_start = header_row + 2

    # Totals: {(entry_type, account): sum_of_abs_amounts}
    type_totals = {}
    acct_totals = {}
    transfer_count = 0
    entry_count = 0

    for row_idx in range(data_start, ws.max_row + 1):
        row = ws[row_idx]
        date_val = row[1].value if len(row) > 1 else None
        desc_val = row[2].value if len(row) > 2 else None

        if not date_val:
            continue
        if isinstance(desc_val, str) and 'C/F' in desc_val.upper():
            continue

        entry_date = parse_date(str(date_val))
        if entry_date is None:
            continue

        for letter, acct, etype in col_map:
            if etype is None:
                continue
            col_idx = openpyxl.utils.column_index_from_string(letter) - 1
            if col_idx >= len(row):
                continue
            val = safe_float(row[col_idx].value)
            if val == 0:
                continue

            abs_val = abs(val)
            key = (etype, acct)
            type_totals[key] = type_totals.get(key, 0) + abs_val
            acct_totals[acct] = acct_totals.get(acct, 0) + abs_val

            if etype == 'bank_transfer':
                transfer_count += 1
            else:
                entry_count += 1

    return type_totals, acct_totals, entry_count, transfer_count


def main():
    print("Loading Excel...")
    wb = load_workbook()

    conn = pymysql.connect(**DB_CONFIG)
    cur = conn.cursor()

    fy_sheets = [s for s in wb.sheetnames if s.startswith('FY')]

    # Collect Excel totals per FY
    excel_by_fy = {}
    excel_grand_type = {}
    excel_grand_acct = {}

    for fy in fy_sheets:
        ws = wb[fy]
        type_totals, acct_totals, entry_count, transfer_count = parse_sheet_totals(ws, fy)
        excel_by_fy[fy] = {
            'type_totals': type_totals,
            'acct_totals': acct_totals,
            'entry_count': entry_count,
            'transfer_count': transfer_count,
        }
        for key, val in type_totals.items():
            excel_grand_type[key] = excel_grand_type.get(key, 0) + val
        for acct, val in acct_totals.items():
            excel_grand_acct[acct] = excel_grand_acct.get(acct, 0) + val

    # Collect DB totals per FY
    db_by_fy = {}
    cur.execute("""
        SELECT financial_year, entry_type,
               CASE
                   WHEN account_id = 1 THEN 'cash'
                   WHEN account_id = 2 THEN 'ab'
                   WHEN account_id = 3 THEN 'sbi'
                   WHEN account_id = 4 THEN 'hdfc'
                   ELSE 'unknown'
               END as acct_key,
               SUM(amount) as total, COUNT(*) as cnt
        FROM entries
        GROUP BY financial_year, entry_type, acct_key
    """)
    for fy, etype, acct_key, total, cnt in cur.fetchall():
        if fy not in db_by_fy:
            db_by_fy[fy] = {'type_totals': {}, 'acct_totals': {}}
        key = (etype, acct_key)
        db_by_fy[fy]['type_totals'][key] = float(total)

    # Also get per-account totals across all FYs
    db_grand_type = {}
    db_grand_acct = {}
    cur.execute("""
        SELECT entry_type,
               CASE
                   WHEN account_id = 1 THEN 'cash'
                   WHEN account_id = 2 THEN 'ab'
                   WHEN account_id = 3 THEN 'sbi'
                   WHEN account_id = 4 THEN 'hdfc'
                   ELSE 'unknown'
               END as acct_key,
               SUM(amount) as total
        FROM entries
        GROUP BY entry_type, acct_key
    """)
    for etype, acct_key, total in cur.fetchall():
        db_grand_type[(etype, acct_key)] = float(total)
        db_grand_acct[acct_key] = db_grand_acct.get(acct_key, 0) + float(total)

    # ===== COMPARISON =====
    print("\n" + "=" * 70)
    print("GRAND TOTAL COMPARISON (all FYs combined)")
    print("=" * 70)

    all_keys = set(list(excel_grand_type.keys()) + list(db_grand_type.keys()))
    max_diff = 0
    max_diff_key = None
    mismatches = []

    for key in sorted(all_keys):
        excel_val = excel_grand_type.get(key, 0)
        db_val = db_grand_type.get(key, 0)
        diff = abs(excel_val - db_val)
        status = "✓" if diff < 0.01 else "✗"
        if diff >= 0.01:
            mismatches.append((key, excel_val, db_val, diff))
            if diff > max_diff:
                max_diff = diff
                max_diff_key = key
        if diff >= 0.01 or True:  # Show all for now
            print(f"  {status} {key[1]:10s} / {key[0]:25s}: Excel=₹{excel_val:>14,.2f}  DB=₹{db_val:>14,.2f}  diff=₹{diff:>12,.2f}")

    print(f"\n  Mismatches: {len(mismatches)}")
    if max_diff_key:
        print(f"  Max diff: ₹{max_diff:,.2f} at {max_diff_key}")

    # Per-account totals
    print("\n" + "=" * 70)
    print("PER-ACCOUNT TOTALS (all FYs combined)")
    print("=" * 70)

    acct_names = {'cash': 'Cash', 'ab': 'Andhra Bank', 'sbi': 'SBI', 'hdfc': 'HDFC'}
    for acct in ['cash', 'ab', 'sbi', 'hdfc']:
        excel_val = excel_grand_acct.get(acct, 0)
        db_val = db_grand_acct.get(acct, 0)
        diff = abs(excel_val - db_val)
        status = "✓" if diff < 0.01 else "✗"
        print(f"  {status} {acct_names[acct]:15s}: Excel=₹{excel_val:>14,.2f}  DB=₹{db_val:>14,.2f}  diff=₹{diff:>12,.2f}")

    # Per-FY comparison (summary)
    print("\n" + "=" * 70)
    print("PER-FY ENTRY COUNTS (non-transfer)")
    print("=" * 70)

    for fy in fy_sheets:
        excel_cnt = excel_by_fy[fy]['entry_count']
        # Get DB count for this FY (non-transfer)
        cur.execute("SELECT COUNT(*) FROM entries WHERE financial_year = %s AND entry_type != 'bank_transfer'", (fy,))
        db_cnt = cur.fetchone()[0]
        diff = abs(excel_cnt - db_cnt)
        status = "✓" if diff == 0 else "✗"
        print(f"  {status} {fy}: Excel={excel_cnt:>4d}  DB={db_cnt:>4d}  diff={diff}")

    # Per-FY transfer counts
    print("\n" + "=" * 70)
    print("PER-FY TRANSFER COUNTS")
    print("=" * 70)

    for fy in fy_sheets:
        excel_cnt = excel_by_fy[fy]['transfer_count']
        cur.execute("SELECT COUNT(*) FROM entries WHERE financial_year = %s AND entry_type = 'bank_transfer'", (fy,))
        db_cnt = cur.fetchone()[0]
        diff = excel_cnt - db_cnt
        status = "✓" if diff == 0 else "~"  # Transfers may differ due to pairing
        print(f"  {status} {fy}: Excel_columns={excel_cnt:>4d}  DB_entries={db_cnt:>4d}  diff={diff:+d}")

    # Per-FY balance check
    print("\n" + "=" * 70)
    print("PER-FY INCOME vs EXPENSE COMPARISON")
    print("=" * 70)

    INCOME_TYPES = ['donation_income', 'bank_interest', 'fd_interest', 'mf_interest', 'rd_interest',
                    'fd_maturity', 'rd_maturity', 'mf_maturity', 'it_return']
    EXPENSE_TYPES = ['donation_expense', 'bank_charges', 'professional_charges', 'stationary_charges',
                     'courier_charges', 'it_paid', 'fd_investment', 'rd_investment', 'mf_investment']

    for fy in fy_sheets:
        # Excel income (non-transfer, non-balance)
        excel_income = sum(v for (etype, acct), v in excel_by_fy[fy]['type_totals'].items()
                          if etype in INCOME_TYPES)
        excel_expense = sum(v for (etype, acct), v in excel_by_fy[fy]['type_totals'].items()
                           if etype in EXPENSE_TYPES)

        # DB income
        placeholders_i = ','.join(['%s'] * len(INCOME_TYPES))
        cur.execute(f"SELECT SUM(amount) FROM entries WHERE financial_year = %s AND entry_type IN ({placeholders_i})",
                    [fy] + INCOME_TYPES)
        db_income = float(cur.fetchone()[0] or 0)

        placeholders_e = ','.join(['%s'] * len(EXPENSE_TYPES))
        cur.execute(f"SELECT SUM(amount) FROM entries WHERE financial_year = %s AND entry_type IN ({placeholders_e})",
                    [fy] + EXPENSE_TYPES)
        db_expense = float(cur.fetchone()[0] or 0)

        inc_diff = abs(excel_income - db_income)
        exp_diff = abs(excel_expense - db_expense)
        inc_status = "✓" if inc_diff < 0.01 else "✗"
        exp_status = "✓" if exp_diff < 0.01 else "✗"

        print(f"  {fy}: Income  {inc_status} Excel=₹{excel_income:>12,.2f}  DB=₹{db_income:>12,.2f}  diff=₹{inc_diff:>10,.2f}")
        print(f"       Expense {exp_status} Excel=₹{excel_expense:>12,.2f}  DB=₹{db_expense:>12,.2f}  diff=₹{exp_diff:>10,.2f}")

    cur.close()
    conn.close()
    print("\nDone!")


if __name__ == '__main__':
    main()
