#!/usr/bin/env python3
"""Extract MASTER-sheet rows from the decrypted Excel file.

Usage:
    python3 extract_master_rows.py --input decrypted.xlsx --output rows.json

Reads the MASTER sheet from the decrypted Excel workbook and writes the cell
values as a JSON array of objects keyed by column letter.
"""

import argparse
import json
import os
import sys

import openpyxl


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True, help="Decrypted Excel file")
    ap.add_argument("--output", required=True, help="Output JSON file")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print(f"ERROR: input file not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    try:
        wb = openpyxl.load_workbook(args.input, data_only=True)
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: cannot load workbook: {exc}", file=sys.stderr)
        sys.exit(1)

    if "MASTER" not in wb.sheetnames:
        print(f"ERROR: MASTER sheet not found. Sheets: {wb.sheetnames}", file=sys.stderr)
        sys.exit(1)

    ws = wb["MASTER"]
    rows = []

    for row in ws.iter_rows(values_only=True):
        cell_values = {}
        for col_idx, value in enumerate(row):
            col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
            cell_values[col_letter] = value
        rows.append(cell_values)

    with open(args.output, "w") as f:
        json.dump(rows, f, default=str)

    print(f"Wrote {len(rows)} rows to {args.output}")


if __name__ == "__main__":
    main()
